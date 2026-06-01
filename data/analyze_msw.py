"""Analyze MassDEP 2024 MSW Survey to extract per-household trash cost signals."""
import openpyxl

wb = openpyxl.load_workbook('data/trash_analysis/2024_MSW_survey.xlsx', data_only=True)
ws = wb['Final']

# Column indices (1-based, from probe above)
COL = {
    'muni': 1,
    'contact': 2,
    'total_hh': 3,
    'hh_served_trash': 4,
    'hh_served_recyc': 5,
    'trash_type': 6,
    'cart_size': 7,
    'recyc_type': 8,
    'recyc_freq': 9,
    'food_type': 11,
    'yard_curb_weeks': 12,
    'fund_property_tax': 19,
    'fund_annual_fee': 20,
    'fund_xfer_fee': 21,
    'fund_pervisit_fee': 22,
    'fund_payt': 23,
    'annual_fee_amt': 24,
    'xfer_fee_amt': 25,
    'pervisit_fee_amt': 26,
    'payt_smart': 27,
    'enforced_limits': 28,
    'max_bags_week': 29,
    'barrel_size': 30,
    'trash_tons': 56,
    'tip_fee': 64,
}

def yesish(v):
    if v is None: return False
    s = str(v).strip().lower()
    return s in ('y','yes','true','1','x','on')

def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace(',', '').replace('$','').strip()
    try: return float(s)
    except ValueError: return None

# Data rows start at row 5
rows = []
for r in range(5, ws.max_row + 1):
    rec = {k: ws.cell(r, c).value for k, c in COL.items()}
    if not rec['muni']: continue
    rows.append(rec)

print(f'Total muni rows: {len(rows)}')

# Inventory funding mechanisms
fund_breakdown = {'property_tax_only': 0, 'annual_fee': 0, 'payt': 0, 'xfer_fee': 0, 'pervisit_fee': 0, 'no_program': 0}
for rec in rows:
    has_pt = yesish(rec['fund_property_tax'])
    has_fee = yesish(rec['fund_annual_fee'])
    has_payt = yesish(rec['fund_payt'])
    has_xfer = yesish(rec['fund_xfer_fee'])
    has_pv = yesish(rec['fund_pervisit_fee'])
    if not any([has_pt, has_fee, has_payt, has_xfer, has_pv]):
        fund_breakdown['no_program'] += 1
        continue
    if has_payt: fund_breakdown['payt'] += 1
    if has_fee: fund_breakdown['annual_fee'] += 1
    if has_xfer: fund_breakdown['xfer_fee'] += 1
    if has_pv: fund_breakdown['pervisit_fee'] += 1
    if has_pt and not (has_fee or has_payt or has_xfer or has_pv):
        fund_breakdown['property_tax_only'] += 1

print('Funding breakdown (non-exclusive except no_program/property_tax_only):')
for k, v in fund_breakdown.items(): print(f'  {k}: {v}')

# Towns with an explicit annual fee — sort descending
print('\n=== TOP TOWNS BY EXPLICIT ANNUAL TRASH FEE ===')
fees = []
for rec in rows:
    fee = num(rec['annual_fee_amt'])
    if fee and fee > 0:
        fees.append((fee, rec['muni'], yesish(rec['fund_payt']), rec['trash_type'], rec['total_hh']))
fees.sort(reverse=True)
print(f'Towns with an annual fee > 0: {len(fees)}')
print(f'{"Rank":<5} {"Town":<25} {"Annual fee":>12} {"PAYT":>5} {"Trash type":<25} {"Total HH":>10}')
for i, (fee, muni, payt, ttype, hh) in enumerate(fees[:30], 1):
    print(f'{i:<5} {muni:<25} ${fee:>10,.0f}   {("Y" if payt else "N"):>5} {(ttype or "")[:25]:<25} {hh or 0:>10}')

# Towns with transfer station access fee
print('\n=== TOWNS WITH TRANSFER STATION ACCESS FEE > 0 ===')
xfer = []
for rec in rows:
    f = num(rec['xfer_fee_amt'])
    if f and f > 0:
        xfer.append((f, rec['muni'], rec['total_hh']))
xfer.sort(reverse=True)
print(f'Count: {len(xfer)}')
for f, m, hh in xfer[:15]:
    print(f'  {m:<25} ${f:>8,.0f}   HH={hh}')

# Trash service type breakdown
from collections import Counter
ttype_ct = Counter(rec['trash_type'] for rec in rows)
print('\n=== TRASH SERVICE TYPE BREAKDOWN ===')
for k, v in sorted(ttype_ct.items(), key=lambda x: -x[1]):
    print(f'  {k}: {v}')

# Save what we found
with open('data/trash_analysis/msw_extracted.csv', 'w') as f:
    f.write('municipality,total_households,hh_served_trash,trash_type,trash_tons,tip_fee_per_ton,fund_property_tax,fund_annual_fee,annual_fee_amt,fund_payt,payt_smart,fund_xfer_fee,xfer_fee_amt\n')
    for rec in rows:
        cols = [
            rec['muni'],
            rec['total_hh'] or '',
            rec['hh_served_trash'] or '',
            (rec['trash_type'] or '').replace(',',';'),
            rec['trash_tons'] or '',
            num(rec['tip_fee']) or '',
            'Y' if yesish(rec['fund_property_tax']) else 'N',
            'Y' if yesish(rec['fund_annual_fee']) else 'N',
            num(rec['annual_fee_amt']) or '',
            'Y' if yesish(rec['fund_payt']) else 'N',
            (rec['payt_smart'] or '').replace(',',';'),
            'Y' if yesish(rec['fund_xfer_fee']) else 'N',
            num(rec['xfer_fee_amt']) or '',
        ]
        f.write(','.join(str(c) for c in cols) + '\n')
print('\nWrote data/trash_analysis/msw_extracted.csv')

# Spot-check Marblehead
mh = [r for r in rows if r['muni'] == 'Marblehead']
if mh:
    print('\n=== Marblehead row ===')
    for k, v in mh[0].items():
        print(f'  {k}: {v}')
